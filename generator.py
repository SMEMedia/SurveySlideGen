from __future__ import annotations

import io
import re
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
for prefix, uri in NS.items():
    if prefix != "rel":
        ET.register_namespace(prefix, uri)


PURCHASE_PROFILES = {
    "Machine Tools": {
        "row_label": "Machine Tools",
        "replacements": [],
    },
    "Manufacturing Software": {
        "row_label": "Software",
        "replacements": [
            ("Machine Tool Buyers", "Manufacturing Software Buyers"),
            ("Machine Tools is the Top Purchase", "Manufacturing Software is a Top Purchase"),
            ("Machine Tool Decision-Makers", "Manufacturing Software Decision-Makers"),
            ("Machine Tool Purchases", "Software Purchases"),
        ],
    },
    "Robotics / Automation": {
        "row_label": "Automation / Robotics",
        "replacements": [
            ("Machine Tool Buyers", "Mfg. Robotics/Automation Buyers"),
            ("Machine Tools is the Top Purchase", "Robotics/Automation is a Top Purchase"),
            ("Machine Tool Decision-Makers", "Robotics & Automation Decision-Makers"),
            ("Machine Tool Purchases", "Robotics/Automation Purchases"),
        ],
    },
    "Machine Tools + Automation": {
        "row_label": "MT & Automation",
        "replacements": [
            ("Machine Tool Buyers", "Machine Tools & Automation Buyers"),
            ("Machine Tools is the Top Purchase", "Manufacturing Equipment is a Top Purchase"),
            ("Machine Tool Decision-Makers", "Machine Tools & Automation Decision-Makers"),
            ("Machine Tool Purchases", "Machine Tools & Automation Purchases"),
        ],
    },
}


@dataclass(frozen=True)
class MetricOption:
    question: str
    label: str
    values: dict[str, float]


def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def read_rows(workbook_file: BinaryIO | bytes) -> list[list[object]]:
    source = io.BytesIO(workbook_file) if isinstance(workbook_file, bytes) else workbook_file
    workbook = load_workbook(source, data_only=True, read_only=True)
    worksheet = workbook.worksheets[0]
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def discover_metric_options(rows: list[list[object]], question_fragment: str) -> list[MetricOption]:
    start = next((i for i, row in enumerate(rows) if question_fragment.lower() in _norm(row[0] if row else "")), -1)
    if start < 0:
        return []
    end = len(rows)
    for index in range(start + 1, len(rows)):
        if re.match(r"^question\s+\d+", str(rows[index][0] or ""), re.I):
            end = index
            break
    header_index = next(
        (i for i in range(start, end) if _norm(rows[i][0]) == "number" and _norm(rows[i][1]) == "choice"),
        -1,
    )
    if header_index < 0:
        return []
    headers = rows[header_index]
    options = []
    for row in rows[header_index + 1 : end]:
        if len(row) < 4 or _norm(row[2]) != "frequency" or row[1] in (None, ""):
            continue
        values = {}
        for index, header in enumerate(headers):
            if index >= len(row) or header in (None, ""):
                continue
            try:
                values[str(header)] = float(row[index])
            except (TypeError, ValueError):
                pass
        options.append(MetricOption(question=question_fragment, label=str(row[1]).strip(), values=values))
    return options


def metric_value(options: list[MetricOption], label: str, column: str = "Total") -> float:
    match = next((item for item in options if _norm(item.label) == _norm(label)), None)
    if match is None:
        raise ValueError(f"Could not find workbook row ‘{label}’.")
    if column not in match.values:
        raise ValueError(f"Could not find workbook column ‘{column}’ for ‘{label}’.")
    return match.values[column]


def _replace_across_nodes(root: ET.Element, old: str, new: str, required: bool = True) -> bool:
    for shape in root.findall(".//p:sp", NS):
        text_nodes = list(shape.findall(".//a:t", NS))
        if not text_nodes:
            continue
        joined = "".join(node.text or "" for node in text_nodes)
        position = joined.find(old)
        if position < 0:
            continue
        before = joined[:position]
        after = joined[position + len(old) :]
        consumed = 0
        inserted = False
        for node in text_nodes:
            text = node.text or ""
            node_start, node_end = consumed, consumed + len(text)
            consumed = node_end
            if node_end <= position or node_start >= position + len(old):
                continue
            prefix = text[: max(0, position - node_start)]
            suffix = text[max(0, position + len(old) - node_start) :]
            if not inserted:
                node.text = prefix + new + suffix
                inserted = True
            else:
                node.text = suffix
        return True
    if required:
        raise ValueError(f"Template text not found: {old}")
    return False


def _replace_shape_text(root: ET.Element, contains: str, replacement: str) -> None:
    for shape in root.findall(".//p:sp", NS):
        nodes = list(shape.findall(".//a:t", NS))
        if not nodes:
            continue
        joined = "".join(node.text or "" for node in nodes)
        if contains.lower() not in joined.lower():
            continue
        lines = replacement.split("\n")
        tx_body = shape.find("p:txBody", NS)
        if tx_body is None:
            raise ValueError(f"Text body not found: {contains}")
        paragraphs = tx_body.findall("a:p", NS)
        if not paragraphs:
            raise ValueError(f"Text paragraphs not found: {contains}")
        while len(paragraphs) < len(lines):
            clone = deepcopy(paragraphs[-1])
            tx_body.append(clone)
            paragraphs.append(clone)
        for paragraph, line in zip(paragraphs, lines):
            paragraph_nodes = paragraph.findall(".//a:t", NS)
            if not paragraph_nodes:
                run = ET.SubElement(paragraph, f"{{{NS['a']}}}r")
                ET.SubElement(run, f"{{{NS['a']}}}rPr", {"lang": "en-US"})
                paragraph_nodes = [ET.SubElement(run, f"{{{NS['a']}}}t")]
            paragraph_nodes[0].text = line
            for node in paragraph_nodes[1:]:
                node.text = ""
        for paragraph in paragraphs[len(lines):]:
            tx_body.remove(paragraph)
        return
    raise ValueError(f"Template shape not found: {contains}")


def _set_chart_cache(xml_bytes: bytes, categories: list[str], values: list[float]) -> bytes:
    root = ET.fromstring(xml_bytes)
    series = root.find(".//c:ser", NS)
    if series is None:
        raise ValueError("Chart series not found")
    cat_ref = series.find(".//c:cat/c:strRef", NS)
    val_ref = series.find(".//c:val/c:numRef", NS)
    if cat_ref is None or val_ref is None:
        raise ValueError("Expected category/value references were not found")
    str_cache = cat_ref.find("c:strCache", NS)
    num_cache = val_ref.find("c:numCache", NS)
    for cache, items in ((str_cache, categories), (num_cache, values)):
        if cache is None:
            continue
        for point in list(cache.findall("c:pt", NS)):
            cache.remove(point)
        count = cache.find("c:ptCount", NS)
        if count is None:
            count = ET.SubElement(cache, f"{{{NS['c']}}}ptCount")
        count.set("val", str(len(items)))
        for index, item in enumerate(items):
            point = ET.SubElement(cache, f"{{{NS['c']}}}pt", {"idx": str(index)})
            ET.SubElement(point, f"{{{NS['c']}}}v").text = str(item)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _update_embedded_workbook(data: bytes, categories: list[str], values: list[float]) -> bytes:
    source = io.BytesIO(data)
    workbook = load_workbook(source)
    worksheet = workbook.worksheets[0]
    for index, (category, value) in enumerate(zip(categories, values), start=2):
        worksheet.cell(index, 1).value = category
        worksheet.cell(index, 2).value = value
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _chart_embedding_name(entries: dict[str, bytes], chart_name: str) -> str | None:
    rel_name = f"ppt/charts/_rels/{Path(chart_name).name}.rels"
    if rel_name not in entries:
        return None
    root = ET.fromstring(entries[rel_name])
    for relationship in root.findall("rel:Relationship", NS):
        target = relationship.attrib.get("Target", "")
        if "embeddings/" in target:
            return "ppt/embeddings/" + target.split("embeddings/", 1)[1]
    return None


def _percent(value: float) -> str:
    return f"{round(value * 100)}%"


def _ratio(value: float, denominator: int) -> str:
    return f"{max(1, round(value * denominator))} in {denominator}"


def generate_presentation(
    *,
    template_bytes: bytes,
    rows: list[list[object]],
    company_name: str,
    report_year: int,
    qualified_value: float,
    qualified_description: str,
    qualified_denominator: int,
    industry_value: float,
    industry_description: str,
    purchase_profile_name: str,
) -> bytes:
    profile = PURCHASE_PROFILES[purchase_profile_name]
    purchase_options = discover_metric_options(rows, "Which types of products/services do you influence or purchase?")
    decision_options = discover_metric_options(rows, "What is your role in purchasing decisions?")
    decision_value = metric_value(decision_options, "purchase influence net", "Total")

    category_labels = ["Other", "Consulting", "Additive Equipment", "Training", "Robotics / Automation", "Software", "Machine Tools"]
    source_labels = ["Other", "Consulting", "Additive Equipment", "Training", "Automation / Robotics", "Software", "Machine Tools"]
    category_values = [metric_value(purchase_options, label, "Total") for label in source_labels]
    platform_columns = ["Video or Podcast", "Mfg Weekly Newsletter", "SME Event Attendees", "Total Magazine Audience", "AM.org USers"]
    platform_labels = ["Podcast/Video Audience", "Mfg. Weekly Newsletter Subscribers", "SME Event Attendees", "MET Magazine Readers", "AM.org Users"]
    platform_values = [metric_value(purchase_options, profile["row_label"], column) for column in platform_columns]

    with zipfile.ZipFile(io.BytesIO(template_bytes), "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}

    slide1 = ET.fromstring(entries["ppt/slides/slide1.xml"])
    _replace_across_nodes(slide1, "Kennametal", company_name)
    _replace_across_nodes(slide1, "Kennametal", company_name)
    _replace_across_nodes(slide1, "Kennametal", company_name)
    _replace_across_nodes(slide1, "35K+", "X")
    _replace_across_nodes(slide1, "2 in 5", _ratio(qualified_value, qualified_denominator))
    _replace_across_nodes(slide1, "43%", _percent(qualified_value))
    _replace_across_nodes(slide1, "Hold Manufacturing Engineering, Production, or C-Suite Leadership Roles at their Organization", qualified_description)
    _replace_across_nodes(slide1, "75%", _percent(industry_value))
    _replace_across_nodes(slide1, "Work across Aerospace & Defense, Automotive, Industrial Machinery & Equipment, Medical Device, & Job Shops", industry_description, required=False)
    _replace_across_nodes(slide1, "8 in 10", _ratio(decision_value, 10))
    _replace_across_nodes(slide1, "82%", _percent(decision_value))
    _replace_across_nodes(slide1, "2026", str(report_year))
    entries["ppt/slides/slide1.xml"] = ET.tostring(slide1, encoding="utf-8", xml_declaration=True)

    slide2 = ET.fromstring(entries["ppt/slides/slide2.xml"])
    for old, new in profile["replacements"]:
        _replace_across_nodes(slide2, old, new)
    _replace_across_nodes(slide2, "2026", str(report_year))
    entries["ppt/slides/slide2.xml"] = ET.tostring(slide2, encoding="utf-8", xml_declaration=True)

    charts = {
        "ppt/charts/chart1.xml": (platform_labels, platform_values),
        "ppt/charts/chart2.xml": (category_labels, category_values),
    }
    for chart_name, (labels, values) in charts.items():
        entries[chart_name] = _set_chart_cache(entries[chart_name], labels, values)
        embedding = _chart_embedding_name(entries, chart_name)
        if embedding and embedding in entries:
            entries[embedding] = _update_embedded_workbook(entries[embedding], labels, values)

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return output.getvalue()
